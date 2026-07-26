#!/usr/bin/env python3
"""Build shareable document deliverables from managed Markdown sources."""

from __future__ import annotations

import argparse
import json
import os
import re
import subprocess
import sys
import tempfile
from dataclasses import dataclass
from pathlib import Path
from typing import Iterable

def preferred_runtime_python() -> Path:
    configured = os.environ.get("CODEX_PRIMARY_RUNTIME_PYTHON")
    if configured:
        return Path(configured).expanduser()
    return Path.home() / ".cache" / "codex-runtimes" / "codex-primary-runtime" / "dependencies" / "python" / "bin" / "python3"


def preferred_runtime_site_packages() -> Path:
    configured = os.environ.get("CODEX_PRIMARY_RUNTIME_SITE_PACKAGES")
    if configured:
        return Path(configured).expanduser()
    return Path.home() / ".cache" / "codex-runtimes" / "codex-primary-runtime" / "dependencies" / "python" / "lib" / "python3.12" / "site-packages"


RUNTIME_PYTHON = preferred_runtime_python()
RUNTIME_SITE_PACKAGES = preferred_runtime_site_packages()

if RUNTIME_PYTHON.exists() and Path(sys.executable).resolve() != RUNTIME_PYTHON.resolve():
    os.execv(str(RUNTIME_PYTHON), [str(RUNTIME_PYTHON), __file__, *sys.argv[1:]])

if RUNTIME_SITE_PACKAGES.exists():
    sys.path.insert(0, str(RUNTIME_SITE_PACKAGES))

from docx import Document  # type: ignore
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement  # type: ignore
from docx.oxml.ns import qn  # type: ignore
from docx.shared import Inches, Pt, RGBColor  # type: ignore


ROOT = Path(__file__).resolve().parent.parent
DOCUMENTS = ROOT / "documents"
DELIVERABLES = DOCUMENTS / "deliverables"

PAGE_WIDTH_IN = 8.5
PAGE_HEIGHT_IN = 11
MARGIN_IN = 1.0
CONTENT_WIDTH_IN = PAGE_WIDTH_IN - (MARGIN_IN * 2)
IMAGE_SUFFIXES = {".png", ".jpg", ".jpeg", ".gif", ".svg", ".webp"}
MERMAID_BLOCK_RE = re.compile(r"```mermaid\s*\n(.*?)\n```", re.S)
OBSIDIAN_EMBED_RE = re.compile(r"!\[\[([^\]]+)\]\]")
OBSIDIAN_LINK_RE = re.compile(r"(?<!!)\[\[([^\]|]+?)(?:\|([^\]]+))?\]\]")
OBSIDIAN_BLOCK_ID_RE = re.compile(r"(\s)\^([A-Za-z0-9_-]+)(?=\s*$)", re.M)


@dataclass
class Section:
    level: int
    title: str
    blocks: list[str]


def parse_frontmatter(text: str) -> tuple[dict[str, str], str]:
    if not text.startswith("---\n"):
        return {}, text
    end = text.find("\n---\n", 4)
    if end < 0:
        return {}, text

    raw = text[4:end]
    body = text[end + 5 :]
    meta: dict[str, str] = {}
    key: str | None = None
    for line in raw.splitlines():
        if re.match(r"^\s+-\s+", line) and key:
            meta[key] = (meta.get(key, "") + "\n" + re.sub(r"^\s+-\s+", "", line)).strip()
            continue
        match = re.match(r"^([A-Za-z_][\w-]*):(?:\s*(.*))?$", line)
        if not match:
            key = None
            continue
        key, value = match.groups()
        meta[key] = (value or "").strip().strip("\"'")
    return meta, body


def publish_body(text: str) -> str:
    if "<!-- publish:start -->" in text and "<!-- publish:end -->" in text:
        return text.split("<!-- publish:start -->", 1)[1].split("<!-- publish:end -->", 1)[0].strip()
    return text


def split_blocks(text: str) -> list[str]:
    lines = [line.rstrip() for line in text.splitlines()]
    blocks: list[str] = []
    current: list[str] = []
    for line in lines:
        if not line.strip():
            if current:
                blocks.append("\n".join(current).strip())
                current = []
            continue
        current.append(line)
    if current:
        blocks.append("\n".join(current).strip())
    return blocks


def parse_sections(text: str) -> list[Section]:
    sections: list[Section] = []
    current = Section(level=0, title="", blocks=[])
    for block in split_blocks(text):
        heading = re.match(r"^(#{1,6})\s+(.*)$", block)
        if heading:
            if current.title or current.blocks:
                sections.append(current)
            current = Section(level=len(heading.group(1)), title=heading.group(2).strip(), blocks=[])
        else:
            current.blocks.append(block)
    if current.title or current.blocks:
        sections.append(current)
    return sections


def clean_inline(text: str) -> str:
    text = re.sub(r"\[\[([^\]|#]+)(?:#[^\]|]+)?\|([^\]]+)\]\]", r"\2", text)
    text = re.sub(r"\[\[([^\]|#]+)(?:#[^\]|]+)?\]\]", lambda m: Path(m.group(1)).stem.replace("-", " "), text)
    text = re.sub(r"\[([^\]]+)\]\(([^)]+)\)", r"\1", text)
    text = re.sub(r"`([^`]+)`", r"\1", text)
    text = re.sub(r"\*\*([^*]+)\*\*", r"\1", text)
    text = re.sub(r"\*([^*]+)\*", r"\1", text)
    return text.strip()


def default_output_for(source: Path, suffix: str) -> Path:
    slug = source.stem
    if source.is_relative_to(DOCUMENTS / "drafts") or source.is_relative_to(DOCUMENTS / "final"):
        return DELIVERABLES / slug / f"{slug}{suffix}"
    raise ValueError(
        "Default output is only available for managed documents under documents/drafts or documents/final; use --output for other Markdown notes."
    )


def display_text_for_target(target: str) -> str:
    return Path(target).stem.replace("-", " ")


def resolve_embed(raw_target: str, source: Path) -> Path | None:
    target = raw_target.split("|", 1)[0].split("#", 1)[0].strip()
    if not target:
        return None
    candidate = (source.parent / target).resolve()
    if candidate.exists():
        return candidate
    root_candidate = (ROOT / target).resolve()
    if root_candidate.exists():
        return root_candidate
    return None


def replace_obsidian_embeds(text: str, source: Path) -> str:
    def repl(match: re.Match[str]) -> str:
        resolved = resolve_embed(match.group(1), source)
        if resolved is None:
            return display_text_for_target(match.group(1))
        if resolved.suffix.lower() in IMAGE_SUFFIXES:
            return f"![](<{resolved.as_posix()}>)"
        return resolved.name

    return OBSIDIAN_EMBED_RE.sub(repl, text)


def replace_obsidian_links(text: str) -> str:
    def repl(match: re.Match[str]) -> str:
        target, alias = match.groups()
        label = alias if alias else display_text_for_target(target)
        if target.startswith("#"):
            if target.startswith("#^ref-"):
                return f"[{label}](#references)"
            anchor = target[1:]
            if anchor.startswith("^"):
                anchor = anchor[1:]
            return f"[{label}](#{anchor})"
        return label

    return OBSIDIAN_LINK_RE.sub(repl, text)


def replace_obsidian_block_ids(text: str) -> str:
    return OBSIDIAN_BLOCK_ID_RE.sub("", text)


def render_mermaid_diagrams(text: str, working: Path) -> str:
    counter = 0
    cached_browsers = sorted(
        Path.home().glob(
            ".cache/puppeteer/chrome-headless-shell/*/chrome-headless-shell-mac-arm64/chrome-headless-shell"
        )
    )
    browser = cached_browsers[-1] if cached_browsers else Path("/Applications/Chromium.app/Contents/MacOS/Chromium")
    puppeteer_config: Path | None = None
    if browser.exists():
        puppeteer_config = working / "puppeteer-config.json"
        puppeteer_config.write_text(
            json.dumps(
                {
                    "executablePath": str(browser),
                    "args": ["--no-sandbox"],
                }
            ),
            encoding="utf-8",
        )

    def repl(match: re.Match[str]) -> str:
        nonlocal counter
        counter += 1
        diagram = match.group(1).strip() + "\n"
        source_path = working / f"diagram-{counter:02d}.mmd"
        image_path = working / f"diagram-{counter:02d}.png"
        source_path.write_text(diagram, encoding="utf-8")
        command = [
            "mmdc",
            "--input",
            str(source_path),
            "--output",
            str(image_path),
            "--theme",
            "neutral",
            "--backgroundColor",
            "transparent",
            "--scale",
            "2",
        ]
        if puppeteer_config is not None:
            command.extend(["--puppeteerConfigFile", str(puppeteer_config)])
        try:
            subprocess.run(command, check=True, capture_output=True, text=True)
        except FileNotFoundError as exc:
            raise RuntimeError("mmdc was not found on PATH. Install mermaid-cli before building PDF deliverables.") from exc
        except subprocess.CalledProcessError as exc:
            stderr = (exc.stderr or "").strip()
            raise RuntimeError(f"Mermaid rendering failed for diagram {counter}: {stderr or exc}") from exc
        return f"![Mermaid diagram {counter}](<{image_path.as_posix()}>)"

    return MERMAID_BLOCK_RE.sub(repl, text)


def strip_named_sections(text: str, headings: list[str]) -> str:
    if not headings:
        return text
    heading_set = {heading.strip().casefold() for heading in headings if heading.strip()}
    if not heading_set:
        return text

    lines = text.splitlines()
    result: list[str] = []
    skip_level: int | None = None

    for line in lines:
        heading = re.match(r"^(#{1,6})\s+(.*)$", line)
        if heading:
            level = len(heading.group(1))
            title = heading.group(2).strip().casefold()
            if skip_level is not None and level <= skip_level:
                skip_level = None
            if skip_level is None and title in heading_set:
                skip_level = level
                continue
        if skip_level is None:
            result.append(line)

    return "\n".join(result).strip()


def preprocess_markdown_for_pdf(source: Path, working: Path, excluded_headings: list[str]) -> Path:
    meta, raw_body = parse_frontmatter(source.read_text(encoding="utf-8"))
    body = publish_body(raw_body).strip()
    body = strip_named_sections(body, excluded_headings)
    body = replace_obsidian_embeds(body, source)
    body = replace_obsidian_links(body)
    body = replace_obsidian_block_ids(body)
    body = render_mermaid_diagrams(body, working)

    if meta.get("title") and not re.search(r"^#\s+", body, re.M):
        body = f"# {meta['title']}\n\n{body}"

    prepared = working / f"{source.stem}.md"
    prepared.write_text(body + "\n", encoding="utf-8")
    return prepared


def build_pdf(source: Path, output: Path, excluded_headings: list[str]) -> Path:
    output.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(prefix="secondbrain-pdf-") as temporary:
        working = Path(temporary)
        prepared = preprocess_markdown_for_pdf(source, working, excluded_headings)
        resource_path = os.pathsep.join([str(source.parent), str(ROOT), str(working)])
        command = [
            "pandoc",
            str(prepared),
            "--from",
            "gfm+pipe_tables",
            "--standalone",
            "--pdf-engine=typst",
            "--resource-path",
            resource_path,
            "--output",
            str(output),
        ]
        try:
            subprocess.run(command, check=True, capture_output=True, text=True)
        except FileNotFoundError as exc:
            raise RuntimeError("pandoc was not found on PATH. Install pandoc before building PDF deliverables.") from exc
        except subprocess.CalledProcessError as exc:
            stderr = (exc.stderr or "").strip()
            raise RuntimeError(f"Pandoc PDF generation failed: {stderr or exc}") from exc
    return output


def list_kind(block: str) -> str | None:
    lines = block.splitlines()
    if all(re.match(r"^- ", line) for line in lines):
        return "bullet"
    if all(re.match(r"^\d+\. ", line) for line in lines):
        return "number"
    return None


def is_markdown_table(block: str) -> bool:
    lines = [line.strip() for line in block.splitlines() if line.strip()]
    if len(lines) < 2:
        return False
    if not all(line.startswith("|") and line.endswith("|") for line in lines[:2]):
        return False
    separator_cells = [cell.strip() for cell in lines[1].strip("|").split("|")]
    return bool(separator_cells) and all(re.match(r"^:?-{3,}:?$", cell) for cell in separator_cells)


def parse_markdown_table(block: str) -> tuple[list[str], list[list[str]]]:
    lines = [line.strip() for line in block.splitlines() if line.strip()]
    rows = [[cell.strip() for cell in line.strip("|").split("|")] for line in lines]
    header = rows[0]
    body = rows[2:]
    width = len(header)
    normalized = [(row + [""] * width)[:width] for row in body]
    return header, normalized


def set_cell_margins(cell, *, top: int = 80, start: int = 80, bottom: int = 80, end: int = 80) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        element = tc_mar.find(qn(f"w:{name}"))
        if element is None:
            element = OxmlElement(f"w:{name}")
            tc_mar.append(element)
        element.set(qn("w:w"), str(value))
        element.set(qn("w:type"), "dxa")


def shade_cell(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def markdown_table_widths(headers: list[str], rows: list[list[str]]) -> list[float]:
    lengths: list[int] = []
    for idx, header in enumerate(headers):
        body_lengths = [len(row[idx]) for row in rows if idx < len(row)]
        lengths.append(max([len(header), *body_lengths, 4]))
    total = sum(lengths) or 1
    raw = [CONTENT_WIDTH_IN * (length / total) for length in lengths]
    minimum = 0.65 if len(headers) >= 5 else 0.9
    adjusted = [max(minimum, width) for width in raw]
    scale = CONTENT_WIDTH_IN / sum(adjusted)
    return [width * scale for width in adjusted]


def set_font(run, name: str, size: float, *, bold: bool = False, italic: bool = False, color: str = "000000") -> None:
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)


def style_paragraph(paragraph, *, after: int, before: int = 0, line: int = 280) -> None:
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line / 240


def configure_styles(doc: Document) -> None:
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)

    style_specs = {
        "Title": ("Calibri", 18, True, "1F1F1F"),
        "Heading 1": ("Calibri", 16, True, "2E74B5"),
        "Heading 2": ("Calibri", 13, True, "2E74B5"),
        "Heading 3": ("Calibri", 12, True, "1F4D78"),
    }
    for name, (font, size, bold, color) in style_specs.items():
        style = doc.styles[name]
        style.font.name = font
        style._element.rPr.rFonts.set(qn("w:ascii"), font)
        style._element.rPr.rFonts.set(qn("w:hAnsi"), font)
        style.font.size = Pt(size)
        style.font.bold = bold
        style.font.color.rgb = RGBColor.from_string(color)

    for missing in ("List Bullet", "List Number"):
        if missing not in [style.name for style in doc.styles]:
            doc.styles.add_style(missing, WD_STYLE_TYPE.PARAGRAPH)


def set_page_geometry(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Inches(PAGE_WIDTH_IN)
    section.page_height = Inches(PAGE_HEIGHT_IN)
    section.top_margin = Inches(MARGIN_IN)
    section.bottom_margin = Inches(MARGIN_IN)
    section.left_margin = Inches(MARGIN_IN)
    section.right_margin = Inches(MARGIN_IN)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)


def add_footer(doc: Document, title: str) -> None:
    section = doc.sections[0]
    footer = section.footer
    para = footer.paragraphs[0]
    para.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    style_paragraph(para, after=0, line=240)
    run = para.add_run(title + " | ")
    set_font(run, "Calibri", 9, color="666666")

    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_separate = OxmlElement("w:fldChar")
    fld_separate.set(qn("w:fldCharType"), "separate")
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")

    page_run = para.add_run()
    page_run._r.append(fld_begin)
    page_run._r.append(instr)
    page_run._r.append(fld_separate)
    page_run._r.append(fld_end)
    set_font(page_run, "Calibri", 9, color="666666")


def add_title_block(doc: Document, title: str, audience: str, updated: str) -> None:
    title_p = doc.add_paragraph(style="Title")
    title_p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    style_paragraph(title_p, after=6, line=240)
    set_font(title_p.add_run(title), "Calibri", 18, bold=True, color="1F1F1F")

    meta = doc.add_paragraph()
    style_paragraph(meta, after=12, line=240)
    meta.alignment = WD_ALIGN_PARAGRAPH.LEFT
    meta_run = meta.add_run(f"Audience: {audience} | Updated: {updated}")
    set_font(meta_run, "Calibri", 10, color="666666")


def add_heading(doc: Document, title: str, level: int) -> None:
    style = "Heading 1" if level <= 2 else "Heading 2" if level == 3 else "Heading 3"
    para = doc.add_paragraph(style=style)
    para.alignment = WD_ALIGN_PARAGRAPH.LEFT
    before = 16 if style == "Heading 1" else 12 if style == "Heading 2" else 8
    after = 8 if style == "Heading 1" else 6 if style == "Heading 2" else 4
    style_paragraph(para, before=before, after=after, line=240)
    run = para.add_run(title)
    if style == "Heading 1":
        set_font(run, "Calibri", 16, bold=True, color="2E74B5")
    elif style == "Heading 2":
        set_font(run, "Calibri", 13, bold=True, color="2E74B5")
    else:
        set_font(run, "Calibri", 12, bold=True, color="1F4D78")


def add_paragraph(doc: Document, text: str) -> None:
    para = doc.add_paragraph()
    style_paragraph(para, after=6, line=264)
    para.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = para.add_run(clean_inline(text))
    set_font(run, "Calibri", 11, color="000000")


def add_list(doc: Document, block: str, ordered: bool) -> None:
    for line in block.splitlines():
        item = re.sub(r"^(- |\d+\. )", "", line).strip()
        para = doc.add_paragraph(style="List Number" if ordered else "List Bullet")
        style_paragraph(para, after=6, line=280)
        para.paragraph_format.left_indent = Inches(0.25)
        para.paragraph_format.first_line_indent = Inches(-0.25)
        run = para.add_run(clean_inline(item))
        set_font(run, "Calibri", 11, color="000000")


def add_markdown_table(doc: Document, block: str) -> None:
    headers, rows = parse_markdown_table(block)
    if not headers:
        return

    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.autofit = False
    widths = markdown_table_widths(headers, rows)
    font_size = 8 if len(headers) >= 5 else 9

    for idx, (cell, header) in enumerate(zip(table.rows[0].cells, headers, strict=True)):
        cell.width = Inches(widths[idx])
        set_cell_margins(cell, top=100, start=100, bottom=100, end=100)
        shade_cell(cell, "D9EAF7")
        para = cell.paragraphs[0]
        style_paragraph(para, after=0, line=240)
        run = para.add_run(clean_inline(header))
        set_font(run, "Calibri", font_size, bold=True, color="1F1F1F")

    for row_values in rows:
        row = table.add_row()
        for idx, cell in enumerate(row.cells):
            cell.width = Inches(widths[idx])
            set_cell_margins(cell, top=90, start=100, bottom=90, end=100)
            para = cell.paragraphs[0]
            style_paragraph(para, after=0, line=240)
            run = para.add_run(clean_inline(row_values[idx]))
            set_font(run, "Calibri", font_size, color="000000")

    spacer = doc.add_paragraph()
    style_paragraph(spacer, after=8, line=240)


def add_summary_table(doc: Document, title: str, summary: str, recommendation: str) -> None:
    table = doc.add_table(rows=3, cols=2)
    table.style = "Table Grid"
    table.autofit = False
    widths = [Inches(1.6), Inches(CONTENT_WIDTH_IN - 1.6)]
    labels = ["Document", "Summary", "Recommendation"]
    values = [title, summary, recommendation]
    for row, label, value in zip(table.rows, labels, values, strict=True):
        for idx, width in enumerate(widths):
            row.cells[idx].width = width
        left, right = row.cells
        left.text = ""
        right.text = ""
        lp = left.paragraphs[0]
        rp = right.paragraphs[0]
        style_paragraph(lp, after=0, line=240)
        style_paragraph(rp, after=0, line=264)
        set_font(lp.add_run(label), "Calibri", 10, bold=True, color="1F1F1F")
        set_font(rp.add_run(value), "Calibri", 10.5, color="000000")


def build_docx(source: Path, output: Path) -> Path:
    meta, raw_body = parse_frontmatter(source.read_text(encoding="utf-8"))
    body = publish_body(raw_body)
    sections = parse_sections(body)
    if sections and sections[0].level == 1:
        sections = sections[1:]

    doc = Document()
    configure_styles(doc)
    set_page_geometry(doc)

    title = meta.get("title", source.stem.replace("-", " ").title())
    audience = meta.get("audience", "")
    updated = meta.get("updated", "")
    add_title_block(doc, title, audience, updated)
    add_footer(doc, title)

    summary_section = next((section for section in sections if section.title.lower() == "summary"), None)
    recommendation = ""
    for section in sections:
        if section.title.lower() == "recommendation":
            recommendation = " ".join(clean_inline(block) for block in section.blocks[:2])
            break
    if summary_section:
        summary = " ".join(clean_inline(block) for block in summary_section.blocks[:2])
        add_summary_table(doc, title, summary, recommendation or "Treat the staged model as a pilot and refine it with evidence.")
        doc.add_paragraph()

    for section in sections:
        lowered = section.title.lower()
        if lowered in {"summary", "notes"}:
            continue
        if lowered == "document" and not section.blocks:
            continue
        if section.title and lowered != "document":
            add_heading(doc, section.title, section.level)
        for block in section.blocks:
            if is_markdown_table(block):
                add_markdown_table(doc, block)
                continue
            kind = list_kind(block)
            if kind == "bullet":
                add_list(doc, block, ordered=False)
            elif kind == "number":
                add_list(doc, block, ordered=True)
            else:
                add_paragraph(doc, block)

    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output)
    return output


def command_docx(args: argparse.Namespace) -> int:
    source = Path(args.source).resolve()
    if not source.exists():
        print(f"ERROR   source not found: {source}")
        return 1
    try:
        output = Path(args.output).resolve() if args.output else default_output_for(source, ".docx")
    except ValueError as exc:
        print(f"ERROR   {exc}")
        return 1
    built = build_docx(source, output)
    print(built)
    return 0


def command_pdf(args: argparse.Namespace) -> int:
    source = Path(args.source).resolve()
    if not source.exists():
        print(f"ERROR   source not found: {source}")
        return 1
    try:
        output = Path(args.output).resolve() if args.output else default_output_for(source, ".pdf")
    except ValueError as exc:
        print(f"ERROR   {exc}")
        return 1
    try:
        built = build_pdf(source, output, args.exclude_heading)
    except RuntimeError as exc:
        print(f"ERROR   {exc}")
        return 1
    print(built)
    return 0


def parser() -> argparse.ArgumentParser:
    p = argparse.ArgumentParser(description=__doc__)
    sub = p.add_subparsers(dest="command", required=True)

    docx = sub.add_parser("docx", help="Build a .docx deliverable from a managed Markdown document")
    docx.add_argument("source", help="Path to the source Markdown document")
    docx.add_argument("--output", help="Optional explicit output path for the .docx")
    docx.set_defaults(func=command_docx)

    pdf = sub.add_parser("pdf", help="Build a .pdf deliverable from a Markdown note or managed document")
    pdf.add_argument("source", help="Path to the source Markdown document")
    pdf.add_argument("--output", help="Optional explicit output path for the .pdf")
    pdf.add_argument(
        "--exclude-heading",
        action="append",
        default=[],
        help="Heading title to omit from the rendered PDF; may be passed more than once",
    )
    pdf.set_defaults(func=command_pdf)
    return p


def main(argv: Iterable[str] | None = None) -> int:
    args = parser().parse_args(argv)
    return args.func(args)


if __name__ == "__main__":
    raise SystemExit(main())
